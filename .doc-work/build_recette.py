from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path
import hashlib, shutil

ROOT = Path(r"D:\USER\WebstormProjects\DuShopPing-By-Assiral-New")
REFERENCE = ROOT / "template-reference.docx"
OUTPUT = ROOT / "Recette-fonctionnelle-DuShopPing.docx"
LOGO = ROOT / "src/assets/images/logo-Dushopping.jpeg"
PHOTOS = [
    ROOT / "src/assets/images/pannier_tisse_1781282354068.jpg",
    ROOT / "src/assets/images/tunique_indigo_1781282369664.jpg",
    ROOT / "src/assets/images/statuettes_vases_1781282386709.jpg",
    ROOT / "src/assets/images/epices_1781282418024.jpg",
]

GREEN = "02603C"; DARK = "152820"; GOLD = "F3C847"; PALE = "E9F5EF"
SLATE = "475569"; LIGHT = "F5F7F6"; RED = "B42318"; BLUE = "2457A6"

expected_hash = "D823CD0115186B34C01C6E4B4DA3BE28B64EE73CAC849DBD62D6F4BB6385B0FB"
assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest().upper() == expected_hash
shutil.copy2(REFERENCE, OUTPUT)
doc = Document(OUTPUT)

# Keep template styles and section properties, replace only body content.
body = doc._element.body
sectPr = body.sectPr
for child in list(body):
    if child is not sectPr:
        body.remove(child)

section = doc.sections[0]
section.top_margin = Inches(0.72); section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.72); section.right_margin = Inches(0.72)
section.header_distance = Inches(0.25); section.footer_distance = Inches(0.3)

def font(run, size=None, bold=None, color=None, name="Georgia"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)
    return run

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn("w:shd"))
    if shd is None: shd = OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc.get_or_add_tcPr(); tcMar = tc.first_child_found_in("w:tcMar")
    if tcMar is None: tcMar = OxmlElement("w:tcMar"); tc.append(tcMar)
    for side, val in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node = tcMar.find(qn("w:"+side))
        if node is None: node=OxmlElement("w:"+side); tcMar.append(node)
        node.set(qn("w:w"), str(val)); node.set(qn("w:type"), "dxa")

def borders(table, color="D8E2DD", size="6"):
    tblPr=table._tbl.tblPr; el=tblPr.find(qn("w:tblBorders"))
    if el is None: el=OxmlElement("w:tblBorders"); tblPr.append(el)
    for edge in ("top","left","bottom","right","insideH","insideV"):
        e=OxmlElement("w:"+edge); e.set(qn("w:val"),"single"); e.set(qn("w:sz"),size); e.set(qn("w:color"),color); el.append(e)

def set_repeat(row):
    trPr=row._tr.get_or_add_trPr(); rep=OxmlElement("w:tblHeader"); rep.set(qn("w:val"),"true"); trPr.append(rep)

def set_width(cell, inches):
    cell.width = Inches(inches)
    tcPr=cell._tc.get_or_add_tcPr(); tcW=tcPr.find(qn("w:tcW"))
    if tcW is None: tcW=OxmlElement("w:tcW"); tcPr.append(tcW)
    tcW.set(qn("w:w"), str(int(inches*1440))); tcW.set(qn("w:type"),"dxa")

def p(text="", style=None, align=None, before=0, after=5, keep=False):
    para=doc.add_paragraph(style=style)
    if text: font(para.add_run(text), 10.2, color=DARK)
    if align is not None: para.alignment=align
    para.paragraph_format.space_before=Pt(before); para.paragraph_format.space_after=Pt(after)
    para.paragraph_format.line_spacing=1.08
    para.paragraph_format.keep_with_next=keep
    return para

def title(text):
    para=doc.add_paragraph(style="Heading 1"); para.paragraph_format.space_before=Pt(10); para.paragraph_format.space_after=Pt(6); para.paragraph_format.keep_with_next=True
    r=para.add_run(text); font(r, 19, True, GREEN)
    return para

def h2(text):
    para=doc.add_paragraph(style="Heading 2"); para.paragraph_format.space_before=Pt(8); para.paragraph_format.space_after=Pt(4); para.paragraph_format.keep_with_next=True
    font(para.add_run(text), 13, True, DARK); return para

def callout(label, text, fill=PALE, accent=GREEN):
    t=doc.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; set_width(t.cell(0,0),7.0); borders(t, accent, "10")
    c=t.cell(0,0); shade(c,fill); margins(c,150,170,150,170)
    pr=c.paragraphs[0]; font(pr.add_run(label.upper()+"\n"),9,True,accent); font(pr.add_run(text),10.2,False,DARK)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def table(headers, rows, widths, header_fill=GREEN):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; borders(t)
    set_repeat(t.rows[0])
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; set_width(c,widths[i]); shade(c,header_fill); margins(c)
        pr=c.paragraphs[0]; font(pr.add_run(h),8.5,True,"FFFFFF"); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            set_width(cells[i],widths[i]); margins(cells[i]);
            if ridx%2: shade(cells[i],"F7FAF8")
            pr=cells[i].paragraphs[0]; font(pr.add_run(str(val)),8.6,False,DARK)
    return t

def checklist(items):
    for item in items:
        para=doc.add_paragraph(); para.paragraph_format.left_indent=Inches(.12); para.paragraph_format.first_line_indent=Inches(-.12); para.paragraph_format.space_after=Pt(3)
        font(para.add_run("[ ] "),10,True,GREEN); font(para.add_run(item),9.5,False,DARK)

def page_break(): doc.add_page_break()

def add_page_number(paragraph):
    paragraph.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    font(paragraph.add_run("DuShopPing  |  Recette fonctionnelle  |  "),8,False,SLATE)
    fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); paragraph._p.append(fld)

# Styles retained from reference, normalized to the reference's dominant Georgia system.
for name, sz, col in [("Normal",10.2,DARK),("Title",32,GREEN),("Subtitle",13,SLATE),("Heading 1",19,GREEN),("Heading 2",13,DARK),("Heading 3",10.5,BLUE)]:
    if name in doc.styles:
        st=doc.styles[name]; st.font.name="Georgia"; st._element.get_or_add_rPr().rFonts.set(qn("w:ascii"),"Georgia"); st._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"),"Georgia"); st.font.size=Pt(sz); st.font.color.rgb=RGBColor.from_string(col)

# Header/footer
header=section.header
for para in header.paragraphs: para._element.getparent().remove(para._element)
hp=header.add_paragraph(); hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(hp.add_run("PROTOCOLE DE RECETTE • ASSIRAL DB"),8,True,GREEN)
footer=section.footer
for para in footer.paragraphs: para._element.getparent().remove(para._element)
add_page_number(footer.add_paragraph())

# Cover
cp=doc.add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.space_after=Pt(18)
cp.add_run().add_picture(str(LOGO), width=Inches(2.6))
band=doc.add_table(rows=1,cols=1); band.alignment=WD_TABLE_ALIGNMENT.CENTER; band.autofit=False; set_width(band.cell(0,0),7); shade(band.cell(0,0),GREEN); borders(band,GREEN,"4"); margins(band.cell(0,0),300,250,300,250)
bp=band.cell(0,0).paragraphs[0]; bp.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(bp.add_run("RECETTE FONCTIONNELLE\n"),27,True,"FFFFFF"); font(bp.add_run("DuShopPing — Marketplace ASSIRAL"),15,False,GOLD)
p("Guide d’exécution destiné au collègue chargé de valider les parcours Acheteur, Vendeur et Administrateur.",align=WD_ALIGN_PARAGRAPH.CENTER,before=18,after=10)
table(["Version","Date","Statut"],[["1.0","31 août 2026","Prêt pour exécution"]],[2.1,2.3,2.6])
p("Comptes de test : fournis séparément par le responsable de recette.",align=WD_ALIGN_PARAGRAPH.CENTER,before=12,after=12)
photo_table=doc.add_table(rows=1,cols=4); photo_table.alignment=WD_TABLE_ALIGNMENT.CENTER; photo_table.autofit=False
for i,img in enumerate(PHOTOS):
    c=photo_table.cell(0,i); set_width(c,1.72); margins(c,40,40,40,40); shade(c,"FFFFFF")
    pr=c.paragraphs[0]; pr.alignment=WD_ALIGN_PARAGRAPH.CENTER; pr.add_run().add_picture(str(img),width=Inches(1.55),height=Inches(1.05))
p("Document interne • Ne pas utiliser de véritables pièces d’identité ou preuves de paiement en environnement de test.",align=WD_ALIGN_PARAGRAPH.CENTER,before=14,after=0)
page_break()

title("1. Contrôle du document")
table(["Champ","Valeur"],[["Nom","Recette fonctionnelle DuShopPing"],["Responsable","____________________________"],["Testeur","____________________________"],["Environnement","[ ] Local   [ ] Préproduction   [ ] Autre : __________"],["URL","____________________________"],["Période d’exécution","Du ____ / ____ / ______ au ____ / ____ / ______"],["Décision finale","[ ] GO   [ ] GO sous réserve   [ ] NO GO"]],[2.0,5.0])
h2("Mode d’emploi")
p("Exécuter les scénarios dans l’ordre proposé. Pour chaque scénario, cocher le résultat, joindre une preuve lorsque demandée et créer une anomalie en cas d’écart. Ne jamais corriger directement la base pendant l’exécution : toute intervention manuelle invalide le scénario concerné.")
callout("Règle de décision", "Un scénario critique échoué entraîne un NO GO jusqu’à correction et nouvelle exécution. Un scénario majeur échoué exige une décision explicite du responsable produit.","FFF5E0","B7791F")
h2("Légende")
table(["Priorité","Définition","Traitement"],[["Critique","Sécurité, paiement, confidentialité, accès ou perte de données","Bloquant"],["Majeure","Parcours principal dégradé ou résultat métier incorrect","Correction avant mise en production"],["Mineure","Défaut visuel ou ergonomique sans perte fonctionnelle","Peut être planifié"]],[1.2,3.4,2.4])

title("2. Objectif et principe de validation")
p("Cette recette vérifie que l’application DuShopPing fonctionne exclusivement avec les données réelles d’ASSIRAL DB, applique les règles de rôle, protège les données privées et assure la continuité des parcours de la création de compte jusqu’à la modération, au Boost et à l’audit.")
callout("Hypothèse de recette", "Si les parcours sont exécutés avec les comptes fournis et sans modification manuelle de la base, alors chaque rôle doit voir uniquement ses fonctions autorisées, chaque action doit produire l’état métier attendu et les événements sensibles doivent être notifiés et audités.")
h2("Critères de succès")
checklist(["100 % des scénarios critiques réussis.","Aucune donnée d’un autre vendeur accessible.","Aucune donnée fictive ou simulation visible.","Les notifications et l’audit correspondent aux actions réellement exécutées.","Les parcours essentiels fonctionnent sur ordinateur et mobile."])

title("3. Cartographie visuelle du parcours")
flow=doc.add_table(rows=1,cols=5); flow.alignment=WD_TABLE_ALIGNMENT.CENTER; flow.autofit=False; borders(flow,GREEN,"8")
steps=[("1","INSCRIPTION","Acheteur ou vendeur"),("2","VÉRIFICATION","Pièces privées"),("3","PUBLICATION","Article + images"),("4","MONÉTISATION","Paiement + Boost"),("5","CONTRÔLE","Notifications + audit")]
for i,(n,lab,desc) in enumerate(steps):
    c=flow.cell(0,i); set_width(c,1.4); shade(c,PALE if i%2==0 else "FFF7DB"); margins(c,170,80,170,80); pr=c.paragraphs[0]; pr.alignment=WD_ALIGN_PARAGRAPH.CENTER
    font(pr.add_run(n+"\n"),19,True,GREEN); font(pr.add_run(lab+"\n"),8.5,True,DARK); font(pr.add_run(desc),7.8,False,SLATE)
p("Le testeur suit ce flux de bout en bout, puis exécute les contrôles négatifs et les vérifications par rôle.",align=WD_ALIGN_PARAGRAPH.CENTER,before=8)
page_break()

title("4. Prérequis et matrice des comptes")
checklist(["Toutes les migrations communiquées ont été exécutées avec succès.","Les variables Supabase, Cloudinary et WhatsApp administrateur sont configurées.","Les comptes Acheteur, Vendeur A, Vendeur B, SuperAdmin, Modérateur et Commercial sont disponibles.","Des fichiers de test non sensibles sont préparés : CNI factice recto/verso, images produit/publicité et preuve de paiement.","Le navigateur autorise les nouveaux onglets WhatsApp et l’accès aux fichiers téléchargés."])
table(["Profil","Identifiant fourni","Mot de passe fourni","État initial"],[["Acheteur","________________","________________","Actif"],["Vendeur A","________________","________________","Boutique en attente"],["Vendeur B","________________","________________","Boutique approuvée"],["SuperAdmin","________________","________________","Actif"],["Modérateur","________________","________________","Actif"],["Commercial","________________","________________","Actif"]],[1.35,2.0,2.0,1.65])
callout("Confidentialité", "Les identifiants sont fournis séparément. Ne pas les recopier dans ce document une fois la recette terminée.","FDECEC",RED)

scenarios = [
 ("AUTH-01","Authentification et séparation des rôles","Critique","Tous","Connexion avec chacun des six comptes.",
  ["Se connecter avec chaque compte fourni.","Vérifier le libellé du profil et la page d’arrivée.","Saisir directement /dashboard puis /admin avec les rôles non autorisés.","Se déconnecter entre chaque essai."],
  ["Chaque compte retrouve son rôle réel.","Acheteur : aucun accès vendeur/admin.","Vendeur : accès /dashboard uniquement.","Administrateurs : menus filtrés selon le rôle.","Les URL interdites affichent Permission insuffisante ou Accès refusé."]),
 ("SELL-01","Profil vendeur et boutique","Majeure","Vendeur A","Compte vendeur actif.",
  ["Ouvrir /dashboard/store.","Modifier nom, localisation et description.","Téléverser un logo puis enregistrer.","Actualiser et se reconnecter."],
  ["Toutes les données persistent.","Le logo apparaît dans l’espace vendeur et le header.","Le vendeur ne peut pas modifier son statut de modération."]),
 ("KYC-01","Soumission des pièces d’identité","Critique","Vendeur A","Fichiers factices JPG/PDF, moins de 5 Mo.",
  ["Ouvrir /dashboard/profile et choisir CNI.","Tenter une soumission sans verso.","Ajouter recto et verso puis transmettre.","Vérifier le statut après actualisation."],
  ["Le verso CNI est obligatoire.","Formats et taille sont contrôlés.","Le dossier passe EnAttente.","La boutique reste masquée pendant l’examen.","Les fichiers ne sont pas publics."]),
 ("KYC-02","Rejet motivé puis approbation","Critique","Modérateur + Vendeur A","Dossier KYC en attente.",
  ["Ouvrir /admin/verifications.","Ouvrir les liens privés recto/verso.","Tenter un rejet sans motif, puis rejeter avec motif.","Vérifier la notification et le motif côté vendeur.","Soumettre de nouvelles pièces et approuver."],
  ["Le rejet sans motif est impossible.","Les liens de document sont temporaires.","Le vendeur reçoit la décision.","Après approbation : vendeur vérifié, boutique approuvée et visible."]),
 ("PROD-01","Publication d’un article","Critique","Vendeur A","Boutique approuvée, quota disponible.",
  ["Ouvrir /dashboard/publish.","Saisir titre, catégorie, prix, description et contacts.","Téléverser plusieurs images puis publier.","Contrôler /dashboard/products, le quota et la fiche publique."],
  ["Redirection vers Mes articles.","Le quota augmente d’une unité.","L’article apparaît sans donnée fictive.","Toutes les images sont navigables.","La fiche utilise la vraie boutique et les vrais contacts."]),
 ("PROD-02","Validation des erreurs de publication","Majeure","Vendeur A","Formulaire de publication ouvert.",
  ["Tester titre vide, prix négatif et description vide.","Tester un fichier trop lourd ou invalide.","Atteindre la limite du quota puis tenter un article supplémentaire."],
  ["Chaque donnée invalide est refusée avec un message clair.","Aucun article partiel n’est créé.","Le dépassement de quota est bloqué côté interface et base."]),
 ("PROD-03","Modification, vente et suppression","Majeure","Vendeur A","Au moins deux articles publiés.",
  ["Modifier titre, prix, statut, contacts et images.","Marquer un article Vendu.","Supprimer un autre article.","Tenter d’ouvrir les anciennes URL publiques."],
  ["Les modifications persistent.","Un article Vendu disparaît du catalogue.","Un article supprimé est introuvable.","Aucun produit d’un autre vendeur ne peut être modifié."]),
 ("MOD-01","Blocage et réactivation d’un article","Critique","Modérateur","Article actif d’une boutique approuvée.",
  ["Ouvrir /admin/products et consulter les détails.","Bloquer l’article.","Contrôler la fiche publique et le compte vendeur.","Réactiver l’article."],
  ["L’article bloqué disparaît publiquement.","Le vendeur ne peut plus le modifier ou le supprimer.","Une notification est reçue.","La réactivation restaure la visibilité si la boutique reste approuvée."]),
 ("SHOP-01","Visibilité de la boutique publique","Critique","Modérateur + visiteur","Boutique approuvée avec produits actifs.",
  ["Ouvrir /seller/{boutiqueId} en navigation privée.","Comparer les informations à l’espace vendeur.","Rejeter la boutique puis actualiser la page et ses produits."],
  ["Seules les informations réelles sont affichées.","Boutique et produits deviennent invisibles après rejet.","Aucune note, photo ou description fictive n’apparaît."]),
 ("BUY-01","Favoris persistants","Majeure","Acheteur","Produits actifs disponibles.",
  ["Ajouter deux articles aux favoris.","Ouvrir le panneau Favoris.","Actualiser puis se reconnecter.","Retirer un article."],
  ["Le compteur et la liste sont corrects.","Les favoris persistent entre sessions.","Un utilisateur ne voit pas ceux d’un autre compte."]),
 ("CONTACT-01","Contact WhatsApp et statistiques","Majeure","Acheteur + Vendeur A","Produit avec numéro WhatsApp.",
  ["Cliquer Appeler le vendeur et Contacter le vendeur.","Vérifier le nouvel onglet et le message prérempli.","Répéter plusieurs fois puis ouvrir /dashboard/stats."],
  ["Le vrai numéro et le vrai titre sont utilisés.","Aucun faux chat interne n’apparaît.","Les interactions sont dédupliquées selon la règle prévue.","Le vendeur ne gonfle pas ses propres statistiques."]),
 ("BOOST-01","Demande de paiement Boost","Critique","Vendeur A","Article actif et plan Boost disponible.",
  ["Ouvrir /dashboard/boost et sélectionner article + plan.","Contrôler #150*47# et le marchand 1044288.","Soumettre référence, opérateur et téléphone facultatif.","Consulter l’historique."],
  ["La demande passe EnAttente.","Le téléphone est clairement facultatif.","Le Commercial reçoit une notification.","La demande est liée au bon article et plan."]),
 ("BOOST-02","Rejet puis validation d’un paiement","Critique","Commercial + Vendeur A","Deux demandes en attente.",
  ["Rejeter la première dans /admin/payments.","Valider la seconde.","Contrôler notifications, historique, campagne et catalogue."],
  ["Rejet : paiement Echoue, aucune campagne.","Validation : paiement Reussi et campagne EnCours.","Le produit devient sponsorisé.","Le quota devient illimité pendant la campagne."]),
 ("ADS-01","Création d’une campagne publicitaire","Majeure","Commercial","Image JPG/PNG/WEBP/GIF de moins de 5 Mo.",
  ["Ouvrir /admin/ads puis Nouvelle campagne.","Choisir un fichier image et vérifier l’aperçu.","Saisir entreprise, emplacement, lien et dates.","Enregistrer, modifier sans changer l’image, puis remplacer l’image."],
  ["Le champ image est un fichier, jamais une URL.","L’aperçu et le téléversement fonctionnent.","L’image existante est conservée en modification.","La campagne respecte emplacement, dates et statut.","Aucune publicité fictive n’apparaît si la base est vide."]),
 ("RBAC-01","Permissions des rôles administratifs","Critique","SuperAdmin, Modérateur, Commercial","Les trois comptes sont disponibles.",
  ["Tester menus et URL directes pour chaque rôle.","Modérateur : tenter paiement/publicité.","Commercial : tenter produit/vendeur.","SuperAdmin : tester quotas, équipe et audit."],
  ["Modérateur : catalogue, vendeurs, vérifications uniquement.","Commercial : Boost, paiements, publicités uniquement.","SuperAdmin : accès complet.","Le serveur refuse aussi les opérations interdites."]),
 ("ADMIN-01","Gestion de l’équipe administrateur","Critique","SuperAdmin","Un utilisateur de test existe.",
  ["Attribuer Moderateur, reconnecter et contrôler.","Changer en Commercial puis reconnecter.","Retirer l’accès.","Tenter de rétrograder ou supprimer le dernier SuperAdmin."],
  ["Les menus suivent le rôle après reconnexion.","Le retrait d’accès est effectif.","Le dernier SuperAdmin est protégé.","Chaque changement est audité."]),
 ("NOTIF-01","Centre de notifications","Majeure","Plusieurs rôles","Actions KYC, modération et paiement disponibles.",
  ["Déclencher une soumission KYC, un blocage et un paiement.","Contrôler compteur, contenu et destinataire.","Ouvrir une notification puis Tout marquer comme lu."],
  ["Chaque notification arrive au bon rôle.","Le lien ouvre la bonne page.","L’état lu persiste après actualisation."]),
 ("AUDIT-01","Journal d’activité exploitable","Critique","SuperAdmin","Plusieurs actions administratives nouvelles.",
  ["Ouvrir /admin/audit.","Filtrer par ressource.","Contrôler auteur, rôle, résumé, motif et date.","Déplier les valeurs avant/après."],
  ["Les actions sont décrites en français.","La ressource, boutique ou vendeur est identifiable.","Le rôle de l’auteur est archivé.","Les mises à jour techniques inutiles ne créent pas de bruit."]),
 ("SEC-01","Isolation des données entre vendeurs","Critique","Vendeur A + Vendeur B","Deux boutiques distinctes.",
  ["Avec Vendeur A, tenter d’accéder aux produits, paiements, statistiques et pièces du Vendeur B.","Modifier manuellement les identifiants dans les requêtes/URL.","Répéter déconnecté."],
  ["Toutes les lectures et modifications privées sont refusées.","Les documents KYC restent privés.","Les RPC administratives refusent les utilisateurs non autorisés."]),
 ("RECOVERY-01","Réinitialisation du mot de passe","Critique","Acheteur","Accès à la boîte e-mail de test.",
  ["Demander le lien via /forgot-password.","Ouvrir /reset-password et tester un mot de passe court.","Définir un mot de passe valide.","Tester ancien puis nouveau mot de passe."],
  ["Le message ne révèle pas l’existence de l’adresse.","Le mot de passe court est refusé.","L’ancien mot de passe échoue et le nouveau fonctionne."]),
 ("RESP-01","Responsive et compatibilité","Majeure","Tous","Chrome, Edge/Firefox et smartphone disponibles.",
  ["Tester 320, 375, 768, 1024 et 1440 px.","Contrôler header, galerie, formulaires, tableaux, menus et fichiers.","Exécuter publication et modération sur mobile."],
  ["Aucun chevauchement ou contenu coupé.","Les tableaux restent consultables.","Les actions principales restent accessibles au clavier et au tactile."]),
 ("NEWS-01","Inscription newsletter","Mineure","Visiteur","Adresse e-mail de test.",
  ["Soumettre une adresse valide.","Soumettre la même adresse une seconde fois.","Tester une adresse invalide."],
  ["L’adresse est enregistrée en minuscules.","Aucun doublon n’est créé.","L’adresse invalide est refusée.","La liste n’est pas publiquement lisible."]),
]

title("5. Synthèse des scénarios")
summary_rows=[[s[0],s[1],s[2],s[3],"[ ] OK  [ ] KO  [ ] N/A"] for s in scenarios]
table(["ID","Scénario","Priorité","Rôle","Résultat"],summary_rows,[0.8,2.8,0.85,1.15,1.4])
page_break()

title("6. Fiches d’exécution détaillées")
p("Une fiche doit être complétée pour chaque scénario. Joindre les captures, références d’anomalie et données de test non sensibles nécessaires à la reproduction.")

for idx,(sid,name,priority,role,precond,steps,expected) in enumerate(scenarios):
    h2(f"{sid} — {name}")
    meta=table(["Priorité","Rôle","Prérequis"],[[priority,role,precond]],[1.0,1.35,4.65], header_fill=DARK)
    p("Procédure",style="Heading 3",before=5,after=2,keep=True)
    rows=[[str(i+1),step,expected[i] if i < len(expected) else "Résultat attendu global ci-dessous", "[ ] OK  [ ] KO"] for i,step in enumerate(steps)]
    table(["#","Action du testeur","Résultat attendu associé","Constat"],rows,[0.35,2.75,2.95,0.95],header_fill=GREEN)
    if len(expected)>len(steps):
        p("Contrôles complémentaires",style="Heading 3",before=5,after=2,keep=True); checklist(expected[len(steps):])
    result=doc.add_table(rows=2,cols=4); result.alignment=WD_TABLE_ALIGNMENT.CENTER; result.autofit=False; borders(result)
    vals=[["Résultat","[ ] Réussi  [ ] Échoué  [ ] Bloqué  [ ] N/A","Anomalie","________________"],["Preuve","Capture / URL / référence : ________________________","Testeur/date","________________"]]
    for r,row in enumerate(vals):
        for c,val in enumerate(row):
            cell=result.cell(r,c); set_width(cell,[0.8,3.0,0.8,2.4][c]); margins(cell); shade(cell,PALE if c%2==0 else "FFFFFF"); font(cell.paragraphs[0].add_run(val),8.3,c%2==0,DARK)
    if idx != len(scenarios)-1: page_break()

page_break(); title("7. Registre des anomalies")
p("Créer une ligne par anomalie. La sévérité doit être évaluée selon l’impact utilisateur et non selon la difficulté technique de correction.")
rows=[[f"ANO-{i:03d}","____ / ____ / ______","________","________________________","[ ] C  [ ] M  [ ] m","[ ] Ouv. [ ] Corr. [ ] Retest","________"] for i in range(1,11)]
table(["ID","Date","Scénario","Résumé","Sévérité","Statut","Responsable"],rows,[0.75,1.0,0.8,1.75,0.9,1.15,0.85])
callout("Preuve minimale", "Toute anomalie doit contenir le scénario, les étapes exactes, le résultat obtenu, le résultat attendu, l’environnement et au moins une capture lorsque le défaut est visuel.","FFF5E0","B7791F")

title("8. Bilan de campagne")
table(["Indicateur","Valeur"],[["Scénarios prévus",str(len(scenarios))],["Scénarios exécutés","_____"],["Réussis","_____"],["Échoués","_____"],["Bloqués","_____"],["Non applicables","_____"],["Anomalies critiques ouvertes","_____"],["Anomalies majeures ouvertes","_____"],["Taux de réussite","_____ %"]],[3.5,3.5])
h2("Limites et risques résiduels")
p("Noter ici tout élément non testé : navigateur indisponible, opérateur Mobile Money non accessible, absence de données suffisantes, problème d’e-mail, service Cloudinary indisponible ou test mobile non réalisé.")
for _ in range(4): p("________________________________________________________________________________",after=3)
h2("Décision")
callout("Conclusion du testeur", "[ ] GO — tous les critères bloquants sont satisfaits\n[ ] GO SOUS RÉSERVE — anomalies acceptées et planifiées\n[ ] NO GO — au moins un risque critique reste ouvert",PALE,GREEN)
table(["Rôle","Nom","Décision / commentaire","Date et visa"],[["Testeur","________________","________________________________","____________"],["Responsable produit","________________","________________________________","____________"],["Responsable technique","________________","________________________________","____________"]],[1.3,1.5,2.9,1.3])

# Document properties and field refresh.
doc.core_properties.title = "Recette fonctionnelle DuShopPing"
doc.core_properties.subject = "Protocole illustré de recette Acheteur, Vendeur et Administrateur"
doc.core_properties.author = "DuShopPing / ASSIRAL"
settings = doc.settings._element
upd = settings.find(qn("w:updateFields"))
if upd is None: upd=OxmlElement("w:updateFields"); settings.append(upd)
upd.set(qn("w:val"),"true")
doc.save(OUTPUT)
print(OUTPUT)
